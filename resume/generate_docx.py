#!/usr/bin/env python3
r"""Generate a Word (.docx) version of a resume variant from its LaTeX source.

The .tex file stays the single source of truth: this script parses the subset of
LaTeX used by resume/*.tex and emits a matching Word document. Edit the .tex and
re-run; never hand-edit the .docx.

Supported constructs (all that the resume variants use):
  \section{Title}
  \resumeSubheading{Title}{Dates}{Company}{}
  \resumeProjectHeading{Name | Descriptor | Link}{}
  \resumeItem{...} and \resumeSubItem{...}
  \textbf{} \emph{} \underline{} \href{}{} and the usual character escapes
  the Technical Skills block (\textbf{Label}{: keywords} \\)

Usage:
  python3 generate_docx.py ai-eval
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parent
FONT = "Cambria"
BODY_SIZE = Pt(10.5)
NAME_SIZE = Pt(19)
CONTACT_SIZE = Pt(9.5)
HEADING_SIZE = Pt(11)
LINE_SPACING = 1.0
HEADING_SPACE_BEFORE = 7
HEADING_SPACE_AFTER = 2
ENTRY_SPACE_BEFORE = 4
BULLET_SPACE_BEFORE = 1
PAGE_WIDTH = Inches(7.5)

ESCAPES = {
    "&": "&", "%": "%", "#": "#", "_": "_", "$": "$", "{": "{", "}": "}",
    "textasciitilde": "~", "textbackslash": "\\",
}

MACRO_NOISE = re.compile(r"\\(Huge|scshape|small|large|Large|bfseries|itshape|normalsize|selectfont)\b")


def strip_comments(source):
    return "\n".join(line for line in source.splitlines() if not line.lstrip().startswith("%"))


def read_group(text, start):
    """Return (content, index_after_closing_brace) for the brace group at start."""
    depth, index = 0, start
    while index < len(text):
        char = text[index]
        if char == "\\":
            index += 2
            continue
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return text[start + 1:index], index + 1
        index += 1
    raise ValueError("unbalanced braces")


def read_macro_args(text, index, count):
    """Read `count` brace groups following index, skipping whitespace."""
    args = []
    for _ in range(count):
        while index < len(text) and text[index].isspace():
            index += 1
        if index >= len(text) or text[index] != "{":
            raise ValueError("expected brace group")
        content, index = read_group(text, index)
        args.append(content)
    return args, index


def unescape(text, strip=True):
    text = text.replace("$\\sim$", "~")
    text = re.sub(r"\\([&%#_${}])", lambda match: ESCAPES.get(match.group(1), match.group(1)), text)
    text = text.replace("---", "\u2014").replace("--", "\u2013")
    text = text.replace("$", "")
    text = MACRO_NOISE.sub("", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip() if strip else text


def parse_runs(text):
    """Split text into styled runs: {text, bold, italic, url}."""
    runs, index, buffer = [], 0, ""

    def flush():
        nonlocal buffer
        if buffer:
            runs.append({"text": unescape(buffer, strip=False), "bold": False, "italic": False, "url": None})
            buffer = ""

    while index < len(text):
        if text[index] == "\\":
            macro = re.match(r"\\([A-Za-z]+)", text[index:])
            if macro and macro.group(1) in ("textbf", "emph", "underline", "href"):
                name = macro.group(1)
                cursor = index + len(macro.group(0))
                args, cursor = read_macro_args(text, cursor, 2 if name == "href" else 1)
                flush()
                if name == "href":
                    label = "".join(run["text"] for run in parse_runs(args[1]))
                    runs.append({
                        "text": label,
                        "bold": False,
                        "italic": False,
                        "url": args[0],
                    })
                else:
                    inner = parse_runs(args[0])
                    for run in inner:
                        if name == "textbf":
                            run["bold"] = True
                        elif name == "emph":
                            run["italic"] = True
                    runs.extend(inner)
                index = cursor
                continue
            escape = re.match(r"\\([&%#_${}]|[A-Za-z]+)", text[index:])
            if escape:
                buffer += escape.group(0)
                index += len(escape.group(0))
                continue
        buffer += text[index]
        index += 1
    flush()
    return [run for run in runs if run["text"]]


def add_hyperlink(paragraph, url, text, size):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    for element in (fonts, color, underline):
        props.append(element)
    size_el = OxmlElement("w:sz")
    size_el.set(qn("w:val"), str(int(size.pt * 2)))
    props.append(size_el)
    run.append(props)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_run(run, size, bold=False, italic=False, small_caps=False):
    run.font.name = FONT
    run.font.size = size
    run.bold = bold
    run.italic = italic
    run.font.small_caps = small_caps


def add_runs(paragraph, runs, size, default_bold=False):
    for run_data in runs:
        if run_data["url"]:
            add_hyperlink(paragraph, run_data["url"], run_data["text"], size)
            continue
        run = paragraph.add_run(run_data["text"])
        style_run(
            run,
            size,
            bold=run_data["bold"] or default_bold,
            italic=run_data["italic"],
        )


def add_bottom_border(paragraph):
    props = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    props.append(borders)


def tight(paragraph, before=0, after=0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = LINE_SPACING
    return paragraph


def add_heading(document, title):
    paragraph = tight(document.add_paragraph(), before=HEADING_SPACE_BEFORE, after=HEADING_SPACE_AFTER)
    run = paragraph.add_run(unescape(title))
    style_run(run, HEADING_SIZE, bold=True, small_caps=True)
    add_bottom_border(paragraph)


def add_entry_title(document, runs):
    paragraph = tight(document.add_paragraph(), before=ENTRY_SPACE_BEFORE)
    paragraph.paragraph_format.tab_stops.add_tab_stop(PAGE_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    add_runs(paragraph, runs, BODY_SIZE)


def add_company(document, text):
    paragraph = tight(document.add_paragraph())
    run = paragraph.add_run(unescape(text))
    style_run(run, BODY_SIZE, italic=True)


def add_bullet(document, text, before=BULLET_SPACE_BEFORE):
    paragraph = tight(document.add_paragraph(style="List Bullet"), before=before)
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    add_runs(paragraph, parse_runs(text), BODY_SIZE)


def add_plain(document, text):
    paragraph = tight(document.add_paragraph())
    add_runs(paragraph, parse_runs(text), BODY_SIZE)


def parse_resume(source):
    body = source.split("\\begin{document}", 1)[1].split("\\end{document}", 1)[0]

    heading_block = body.split("\\section{", 1)[0]
    name = re.search(r"\\textbf\{\\Huge \\scshape ([^}]*)\}", heading_block).group(1).strip()
    contact = " ".join(
        line.strip()
        for line in heading_block.splitlines()
        if line.strip().startswith(("\\small", "\\href"))
    )

    sections = []
    position = 0
    while True:
        match = re.search(r"\\section\{([^}]*)\}", body[position:])
        if not match:
            break
        start = position + match.end()
        next_match = re.search(r"\\section\{", body[start:])
        end = start + (next_match.start() if next_match else len(body) - start)
        sections.append((match.group(1), body[start:end]))
        position = start

    parsed = {"name": name, "contact": contact, "sections": []}
    for title, content in sections:
        if title.startswith("Technical Skills"):
            skills = []
            for line in content.splitlines():
                if "\\textbf{" not in line:
                    continue
                label, cursor = read_macro_args(line, line.index("\\textbf{") + len("\\textbf"), 2)
                skills.append((unescape(label[0]), unescape(label[1]).lstrip(": ").strip()))
            parsed["sections"].append({"title": title, "kind": "skills", "items": skills})
            continue

        entries, index = [], 0
        current = None
        while index < len(content):
            if content.startswith("\\resumeSubheading", index):
                args, index = read_macro_args(content, index + len("\\resumeSubheading"), 4)
                current = {"kind": "role", "title": args[0], "dates": args[1], "company": args[2], "bullets": []}
                entries.append(current)
                continue
            if content.startswith("\\resumeProjectHeading", index):
                args, index = read_macro_args(content, index + len("\\resumeProjectHeading"), 2)
                current = {"kind": "project", "heading": args[0], "bullets": []}
                entries.append(current)
                continue
            if content.startswith("\\resumeItem{", index) or content.startswith("\\resumeSubItem{", index):
                macro = "\\resumeItem" if content.startswith("\\resumeItem{", index) else "\\resumeSubItem"
                args, index = read_macro_args(content, index + len(macro), 1)
                if current is not None:
                    current["bullets"].append({"text": args[0], "plain": macro == "\\resumeSubItem"})
                continue
            index += 1
        parsed["sections"].append({"title": title, "kind": "entries", "items": entries})
    return parsed


def build_docx(parsed, target):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, margin, Inches(0.5))

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = LINE_SPACING

    paragraph = tight(document.add_paragraph(), after=1)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(parsed["name"])
    style_run(run, NAME_SIZE, bold=True, small_caps=True)

    paragraph = tight(document.add_paragraph(), after=2)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(paragraph, parse_runs(parsed["contact"]), CONTACT_SIZE)

    for entry in parsed["sections"]:
        add_heading(document, entry["title"])
        if entry["kind"] == "skills":
            for label, keywords in entry["items"]:
                paragraph = tight(document.add_paragraph())
                paragraph.paragraph_format.left_indent = Inches(0.15)
                paragraph.paragraph_format.first_line_indent = Inches(-0.15)
                label_run = paragraph.add_run(label + ": ")
                style_run(label_run, BODY_SIZE, bold=True)
                text_run = paragraph.add_run(keywords)
                style_run(text_run, BODY_SIZE)
            continue
        for item in entry["items"]:
            if item["kind"] == "role":
                add_entry_title(document, parse_runs(item["title"]) + [{"text": "\t" + unescape(item["dates"]), "bold": False, "italic": False, "url": None}])
                add_company(document, item["company"])
            else:
                add_entry_title(document, parse_runs(item["heading"]))
            for bullet in item["bullets"]:
                if bullet["plain"]:
                    add_plain(document, bullet["text"])
                else:
                    add_bullet(document, bullet["text"])

    document.save(target)
    return target


def main():
    variant = sys.argv[1] if len(sys.argv) > 1 else "ai-eval"
    source_path = BASE / f"{variant}.tex"
    if not source_path.exists():
        sys.exit(f"no such variant: {source_path}")

    parsed = parse_resume(strip_comments(source_path.read_text()))
    target = build_docx(parsed, BASE / f"{variant}.docx")
    print(f"Wrote {target}")
    print(f"  sections: {[entry['title'] for entry in parsed['sections']]}")
    print(f"  entries:  {[len(entry['items']) for entry in parsed['sections'] if entry['kind'] == 'entries']}")


if __name__ == "__main__":
    main()
